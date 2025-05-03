import tkinter as tk
from tkinter import messagebox
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from datetime import datetime

# ========== Validation Function ==========
def validate_inputs():
    name = name_entry.get()
    grade_str = grade_entry.get()

    if not name or not grade_str:
        messagebox.showerror("Input Error", "All fields are required!")
        return False

    try:
        grade = int(grade_entry.get())
    except ValueError:
        messagebox.showerror("Input Error", "Grade must be a valid number.")
        return False

    if grade >= 75:
        print("Passed")
    elif grade <= 74:
        print("Failed")
    else:
        print("Invalid Grade") 

    return True

def grade_letter(grade):
    if grade >= 95:
        return "Passed"
    elif grade >= 90:
        return "Passed"
    elif grade >= 85:
        return "Passed"
    elif grade >= 80:
        return "Passed"
    elif grade >= 75:
        return "Passed"
    else:
        return "Failed"


def save_to_excel():

    if not validate_inputs():
        return

    name = name_entry.get()
    grade = int(grade_entry.get())

    wb = load_workbook("gradinglist.xlsx")
    ws = wb["Score-Tracker-Data"]
    ws.append([name, grade, grade_letter(grade)])
    wb.save("gradinglist.xlsx")

    format_excel()
    messagebox.showinfo("Success", "Data saved successfully!")

    name_entry.delete(0, tk.END)
    grade_entry.delete(0, tk.END)

    

    # ========== Format Excel Sheet ==========
def format_excel():
    wb = load_workbook("gradinglist.xlsx")
    ws = wb["Score-Tracker-Data"]

    # Bold header
    for cell in ws[1]:
        cell.font = Font(bold=True)

    # Adjust column widths
    for col in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max_length + 2

    wb.save("gradinglist.xlsx")

    # ========== View Excel Data ==========
def show_data():
    wb = load_workbook("gradinglist.xlsx")
    ws = wb["Score-Tracker-Data"]

    data_window = tk.Toplevel(window)
    data_window.geometry("1366x768")
    data_window.title("Stored-User-Data")

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        for j, value in enumerate(row):
            label = tk.Label(data_window, text=value, borderwidth=1, relief="solid", padx=6, pady=3)
            label.grid(row=i, column=j)


#====================UI=========
window = tk.Tk()
window.title("Score Tracker")
window.geometry("1366x768")


# Labels
tk.Label(window, text="Name").grid(row=0, column=0, padx=15, pady=10, sticky="nesw")
tk.Label(window, text="Grade").grid(row=1, column=0, padx=15, pady=10, sticky="nesw")


# Entry Fields
name_entry = tk.Entry(window, width=50)
grade_entry = tk.Entry(window, width=50)


name_entry.grid(row=0, column=1, pady=5)
grade_entry.grid(row=1, column=1, pady=5)


# Buttons
tk.Button(window, text="Submit", command=save_to_excel, width=20, bg="#61AD4E", fg="white").grid(row=3, column=0, columnspan=2, pady=10)
tk.Button(window, text="View Stored Data", command=show_data, width=20, bg="#226710", fg="white").grid(row=4, column=0, columnspan=2)

window.mainloop()